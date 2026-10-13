import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# 1) webscraper.io data saaf karna
w = pd.read_csv("webscraper_products.csv")
w["currency"] = w["price"].str.extract(r"^(\D+)")[0].str.strip()
w["price"] = w["price"].str.replace(r"[^\d.]", "", regex=True).astype(float)
w["reviews"] = w["reviews"].str.extract(r"(\d+)")[0].astype(int)

# 2) baqi do sites ka data
s = pd.read_csv("scrapeme_products.csv")
t = pd.read_csv("third_products.csv")

# 3) teeno ko ek table mein jorna
columns = ["source", "category", "name", "price", "currency",
           "rating", "reviews", "description", "image", "link"]
df = pd.concat([w, s, t], ignore_index=True)
df = df.reindex(columns=columns)
df = df.drop_duplicates(subset="link").reset_index(drop=True)

# 4) CSV aur Excel dono save karna
df.to_csv("final_products.csv", index=False, encoding="utf-8-sig")

with pd.ExcelWriter("final_products.xlsx", engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Products")
    ws = writer.sheets["Products"]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        width = min(max(df[col].astype(str).str.len().max(), len(col)) + 2, 50)
        ws.column_dimensions[get_column_letter(i)].width = width

print("Total rows:", len(df))
print(df["source"].value_counts())