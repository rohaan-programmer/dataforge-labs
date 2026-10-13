import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

df = pd.read_csv("final_products.csv")

sample = pd.concat([
    df[df["source"] == "webscraper.io"].head(10),
    df[df["source"] == "scrapeme.live"].head(10),
    df[df["source"] == "scrapingcourse.com"].head(10),
]).reset_index(drop=True)

with pd.ExcelWriter("sample_30_rows.xlsx", engine="openpyxl") as writer:
    sample.to_excel(writer, index=False, sheet_name="Sample")
    ws = writer.sheets["Sample"]
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(sample.columns, start=1):
        width = min(max(sample[col].astype(str).str.len().max(), len(col)) + 2, 40)
        ws.column_dimensions[get_column_letter(i)].width = width

print("Sample rows:", len(sample))